"""折衷デッキのエクスポート（HTML組み立て・PNG/PDF/PPTX/サムネイル生成）"""

import html as html_escape
import io
import json
import re
import subprocess
import tempfile
from pathlib import Path

_ASSET_DIR = Path(__file__).parent.parent / "decks" / "eclectic"
_RENDER_SCRIPT = Path(__file__).parent / "render_deck.mjs"

RENDER_TIMEOUT = 300  # Chromiumレンダリングのタイムアウト（秒）


def is_deck_source(source: str) -> bool:
    """スライドソースが折衷デッキ（HTML）かどうかを判定"""
    stripped = source.lstrip()
    # HTMLコメントをスキップして先頭要素を判定
    stripped = re.sub(r'^(<!--.*?-->\s*)+', '', stripped, flags=re.DOTALL)
    return stripped.startswith('<')


def extract_deck_title(slides_html: str) -> str | None:
    """デッキHTMLからタイトル（最初のh1/h2のテキスト）を抽出"""
    match = re.search(r'<h[12][^>]*>(.*?)</h[12]>', slides_html, re.DOTALL)
    if not match:
        return None
    text = re.sub(r'<br\s*/?>', ' ', match.group(1))
    text = re.sub(r'<[^>]+>', '', text)
    text = html_escape.unescape(text).strip()
    return text or None


def assemble_deck_html(slides_html: str, title: str | None = None, subtitle: str = "") -> str:
    """スライドブロックをシェルに埋め込み、CSS/JSをインライン化した単一HTMLを返す"""
    shell = (_ASSET_DIR / "deck-shell.html").read_text(encoding="utf-8")
    css = (_ASSET_DIR / "deck.css").read_text(encoding="utf-8")
    js = (_ASSET_DIR / "deck.js").read_text(encoding="utf-8")

    resolved_title = title or extract_deck_title(slides_html) or "スライド"
    html = shell.replace("{{DECK_TITLE}}", html_escape.escape(resolved_title))
    html = html.replace("{{DECK_SUBTITLE}}", html_escape.escape(subtitle))

    # 外部参照をインライン化して自己完結型の1ファイルにする（共有・レンダリング両用）
    html = html.replace('<link rel="stylesheet" href="deck.css">', f"<style>\n{css}\n</style>")
    html = html.replace('<script src="deck.js"></script>', f"<script>\n{js}\n</script>")

    start_marker = "<!-- SLIDES:START -->"
    end_marker = "<!-- SLIDES:END -->"
    start = html.index(start_marker) + len(start_marker)
    end = html.index(end_marker)
    return html[:start] + "\n" + slides_html + "\n" + html[end:]


def _render_deck_pngs(slides_html: str, scale: int = 2, limit: int | None = None) -> tuple[list[Path], list[dict]]:
    """デッキをChromiumでレンダリングし、PNGパスのリストとノート情報を返す"""
    tmpdir = Path(tempfile.mkdtemp())
    deck_path = tmpdir / "deck.html"
    deck_path.write_text(assemble_deck_html(slides_html), encoding="utf-8")
    out_dir = tmpdir / "frames"

    cmd = ["node", str(_RENDER_SCRIPT), str(deck_path), str(out_dir), "--scale", str(scale)]
    if limit is not None:
        cmd.extend(["--limit", str(limit)])

    result = subprocess.run(cmd, capture_output=True, text=True, timeout=RENDER_TIMEOUT)
    if result.returncode != 0:
        raise RuntimeError(f"Deck render error: {result.stderr or result.stdout}")

    pngs = sorted(out_dir.glob("slide-*.png"))
    if not pngs:
        raise RuntimeError("Deck render failed: no PNG files created")

    notes_path = out_dir / "notes.json"
    notes = json.loads(notes_path.read_text(encoding="utf-8")) if notes_path.exists() else []
    return pngs, notes


def generate_deck_pptx(slides_html: str, title: str | None = None) -> bytes:
    """折衷デッキをPPTXに変換（全面画像スライド＋スピーカーノート）"""
    from pptx import Presentation
    from pptx.util import Inches

    pngs, notes = _render_deck_pngs(slides_html)
    notes_by_index = {int(item["index"]): item for item in notes}

    prs = Presentation()
    # 1920×1080のデザインに合わせた16:9キャンバス
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, png in enumerate(pngs, start=1):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
        meta = notes_by_index.get(i)
        if meta and meta.get("notes"):
            tf = slide.notes_slide.notes_text_frame
            label = meta.get("label", "")
            tf.text = (f"[{label}] " if label else "") + meta["notes"]

    prs.core_properties.title = title or extract_deck_title(slides_html) or "スライド"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def generate_deck_pdf(slides_html: str) -> bytes:
    """折衷デッキをPDFに変換（レンダリング済みPNGを結合）"""
    import img2pdf

    pngs, _ = _render_deck_pngs(slides_html)
    return img2pdf.convert([str(p) for p in pngs])


def generate_deck_thumbnail(slides_html: str) -> bytes:
    """折衷デッキの1枚目をPNGで生成（OGP用サムネイル）"""
    pngs, _ = _render_deck_pngs(slides_html, scale=1, limit=1)
    return pngs[0].read_bytes()


def generate_deck_standalone_html(slides_html: str, title: str | None = None) -> str:
    """折衷デッキのスタンドアロンHTMLを生成（共有用）"""
    return assemble_deck_html(slides_html, title=title)
